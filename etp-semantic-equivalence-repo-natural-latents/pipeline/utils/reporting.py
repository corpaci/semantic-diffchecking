"""
Writing results: one Excel workbook (+ CSV copies) and confusion-matrix images.

Workbook sheets
  Summary            one row per method, ranked by balanced accuracy: accuracy, balanced
                     accuracy, macro precision / recall / specificity / F1 (mean +- std over
                     splits), weighted F1, Cohen's kappa, MCC, symmetric-method check
  PerClass           method x class: support, TP/FP/FN/TN, precision, recall, specificity,
                     F1, accuracy (pooled over splits, plus mean +- std over splits)
  ConfusionMatrices  every method's pooled 4x4 confusion matrix (counts and row %)
  PerSplit           every metric for every method and split
  Methods            catalogue: family, type, representations used, description
  Config             run settings and data statistics
"""
import os
import re

import numpy as np
import pandas as pd

from .data import LABEL_NAMES
from .metrics import OVERALL_KEYS, PER_CLASS_KEYS

PRETTY = {
    'accuracy': 'Accuracy', 'balanced_accuracy': 'Balanced accuracy',
    'macro_precision': 'Macro precision', 'macro_recall': 'Macro recall',
    'macro_specificity': 'Macro specificity', 'macro_f1': 'Macro F1',
    'weighted_f1': 'Weighted F1', 'cohen_kappa': "Cohen's kappa", 'mcc': 'MCC',
}


def _slug(s):
    return re.sub(r'[^A-Za-z0-9]+', '_', s).strip('_')[:90]


def build_tables(results, ctx):
    ranked = sorted(results, key=lambda r: -r['overall']['balanced_accuracy'])
    summary, per_class, per_split, conf_long, catalogue = [], [], [], [], []
    for rank, r in enumerate(ranked, 1):
        o = r['overall']
        row = {'Rank': rank, 'Family': r['family'], 'Method': r['name'],
               'Type': r['kind'], 'Classifier': r['classifier'],
               'Representations': r['representations'], '#Features': r['n_features'] or None}
        for k in OVERALL_KEYS:
            row[PRETTY[k]] = o[k]
            row[PRETTY[k] + ' std'] = o[k + '_std']
        row['Symmetric method'] = 'yes' if r['symmetric'] else 'no'
        row['Stronger row == Weaker row'] = 'IDENTICAL' if r['stronger_row_equals_weaker_row'] else 'different'
        row['Runtime (s)'] = round(r['runtime_s'], 1)
        summary.append(row)
        for c in r['per_class']:
            pc = {'Rank': rank, 'Family': r['family'], 'Method': r['name'], 'Class': c['class'],
                  'Support': c['support'], 'TP': c['TP'], 'FP': c['FP'], 'FN': c['FN'], 'TN': c['TN']}
            for k in PER_CLASS_KEYS:
                pc[k.capitalize()] = c[k]
            for k in PER_CLASS_KEYS:
                pc[f'{k.capitalize()} mean over splits'] = c[k + '_mean_over_splits']
                pc[f'{k.capitalize()} std over splits'] = c[k + '_std_over_splits']
            per_class.append(pc)
        for s, o_s in zip(ctx.splits, r['per_split']):
            per_split.append({'Method': r['name'], 'Family': r['family'], 'Split seed': s.seed,
                              'Test pairs': int(s.test_pairs.sum()),
                              **{PRETTY[k]: o_s[k] for k in OVERALL_KEYS}})
        cm = r['confusion_pooled']
        for a, ta in enumerate(LABEL_NAMES):
            for b, pb in enumerate(LABEL_NAMES):
                conf_long.append({'Method': r['name'], 'True': ta, 'Predicted': pb, 'Count': int(cm[a, b])})
        catalogue.append({'Method': r['name'], 'Family': r['family'], 'Type': r['kind'],
                          'Representations': r['representations'], 'Description': r['description']})
    return ranked, {
        'Summary': pd.DataFrame(summary),
        'PerClass': pd.DataFrame(per_class),
        'PerSplit': pd.DataFrame(per_split),
        'ConfusionLong': pd.DataFrame(conf_long),
        'Methods': pd.DataFrame(catalogue),
    }


def write_workbook(path, ranked, tables, config_rows):
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill('solid', fgColor='DDE4EE')
    bold = Font(bold=True)
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        for name in ('Summary', 'PerClass', 'PerSplit', 'Methods'):
            tables[name].to_excel(xw, sheet_name=name, index=False)
        pd.DataFrame(config_rows, columns=['Setting', 'Value']).to_excel(xw, sheet_name='Config', index=False)

        # ConfusionMatrices: one block per method (counts | row percentages)
        ws = xw.book.create_sheet('ConfusionMatrices', 2)
        row = 1
        for rank, r in enumerate(ranked, 1):
            cm = r['confusion_pooled']
            pct = cm / np.maximum(cm.sum(axis=1, keepdims=True), 1) * 100
            ws.cell(row, 1, f"#{rank}  {r['name']}").font = bold
            ws.cell(row, 8, f"balanced acc {r['overall']['balanced_accuracy']:.3f} | "
                            f"acc {r['overall']['accuracy']:.3f} | pooled over {len(r['per_split'])} splits")
            row += 1
            ws.cell(row, 1, 'true \\ predicted (counts)').font = bold
            ws.cell(row, 7, 'true \\ predicted (row %)').font = bold
            for b, pb in enumerate(LABEL_NAMES):
                ws.cell(row, 2 + b, pb).font = bold
                ws.cell(row, 8 + b, pb).font = bold
            for a, ta in enumerate(LABEL_NAMES):
                ws.cell(row + 1 + a, 1, ta).font = bold
                ws.cell(row + 1 + a, 7, ta).font = bold
                for b in range(4):
                    ws.cell(row + 1 + a, 2 + b, int(cm[a, b]))
                    c = ws.cell(row + 1 + a, 8 + b, round(float(pct[a, b]), 1))
                    if a == b:
                        c.font = bold
            row += 7
        ws.column_dimensions['A'].width = 26
        ws.column_dimensions['G'].width = 26

        for name in ('Summary', 'PerClass', 'PerSplit', 'Methods', 'Config'):
            ws = xw.sheets[name]
            ws.freeze_panes = 'D2' if name in ('Summary', 'PerClass') else 'B2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font, cell.fill = bold, header_fill
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            df = tables[name] if name in tables else None
            for j, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
                title = str(col[0].value)
                width = 60 if title in ('Method', 'Description', 'Classifier', 'Value') else \
                    (18 if title in ('Family', 'Representations', 'Class', 'Setting') else 12)
                ws.column_dimensions[get_column_letter(j)].width = width
                if df is not None and title in df and df[title].dtype.kind == 'f':
                    for c in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        c[0].number_format = '0.000'
            if name == 'Summary':
                cols = {c.value: c.column_letter for c in ws[1]}
                for key in ('Balanced accuracy', 'Accuracy', 'Macro F1'):
                    L = cols[key]
                    ws.conditional_formatting.add(
                        f'{L}2:{L}{ws.max_row}',
                        ColorScaleRule(start_type='num', start_value=0.25, start_color='F8696B',
                                       mid_type='num', mid_value=0.6, mid_color='FFEB84',
                                       end_type='num', end_value=1.0, end_color='63BE7B'))


def save_confusion_pngs(ranked, out_dir):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    os.makedirs(out_dir, exist_ok=True)
    short = ['equiv', 'stronger', 'weaker', 'incomp']
    for rank, r in enumerate(ranked, 1):
        cm = r['confusion_pooled']
        pct = cm / np.maximum(cm.sum(axis=1, keepdims=True), 1)
        fig, ax = plt.subplots(figsize=(4.2, 3.8))
        ax.imshow(pct, cmap='Blues', vmin=0, vmax=1)
        for a in range(4):
            for b in range(4):
                ax.text(b, a, f'{cm[a, b]}\n{pct[a, b] * 100:.0f}%', ha='center', va='center', fontsize=7,
                        color='white' if pct[a, b] > 0.6 else 'black')
        ax.set_xticks(range(4), short, fontsize=7)
        ax.set_yticks(range(4), short, fontsize=7)
        ax.set_xlabel('predicted', fontsize=8)
        ax.set_ylabel('true', fontsize=8)
        title = r['name'] if len(r['name']) < 60 else r['name'][:57] + '...'
        ax.set_title(f"#{rank} {title}\nbal acc {r['overall']['balanced_accuracy']:.3f}", fontsize=7)
        fig.tight_layout()
        fig.savefig(os.path.join(out_dir, f'{rank:03d}_{_slug(r["name"])}.png'), dpi=110)
        plt.close(fig)


def write_all(results, ctx, config_rows):
    out = ctx.cfg.results_dir
    os.makedirs(out, exist_ok=True)
    ranked, tables = build_tables(results, ctx)
    xlsx = os.path.join(out, 'method_comparison.xlsx')
    write_workbook(xlsx, ranked, tables, config_rows)
    for name in ('Summary', 'PerClass', 'PerSplit', 'ConfusionLong', 'Methods'):
        tables[name].to_csv(os.path.join(out, f'{name.lower()}.csv'), index=False)
    if ctx.cfg.save_confusion_pngs:
        save_confusion_pngs(ranked, os.path.join(out, 'confusion_matrices'))
    return xlsx, tables
